import os, json, re, time, threading, datetime, io, csv, urllib.parse
import hashlib, hmac, base64, collections
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "data")
PUBLIC = os.path.join(BASE, "public")
os.makedirs(DATA, exist_ok=True)

_cfgp = os.path.join(BASE, "config.json")
CFG = json.load(open(_cfgp, encoding="utf-8")) if os.path.exists(_cfgp) else {}
TOKEN = os.environ.get("LIVECHAT_TOKEN") or CFG.get("token", "")
PH = datetime.timezone(datetime.timedelta(hours=int(os.environ.get("PH_TZ_HOURS", CFG.get("ph_tz_hours", 8)))))
PORT = int(os.environ.get("PORT", CFG.get("port", 8787)))
HOST = os.environ.get("HOST") or ("0.0.0.0" if os.environ.get("PORT") else CFG.get("host", "127.0.0.1"))
API = "https://api.livechatinc.com/v3.5/agent/action/list_archives"
CFG_API = "https://api.livechatinc.com/v3.5/configuration/action/"
HDR = {"Authorization": "Basic " + TOKEN, "Content-Type": "application/json"}

PROFANITY = ["putang", "puta", "gago", "gaga", "tanga", "bobo", "ulol", "hayop", "hayup",
             "lintik", "tarantado", "punyeta", "kingina", "kupal", "inutil", "buraot",
             "sakim", "lintek", "burat", "pokpok", "ina mo", "fuck", "shit", "bitch",
             "asshole", "bastard", "stupid", "idiot", "damn", "wtf", "dumbass", "moron"]
SIGNAL_WORDS = ["scam", "fraud", "refund", "magnanakaw", "steal", "rob"]
HOLD_PAT = ["please wait", "for a while", "wait for", "submitted to the relevant",
            "escalat", "still working", "be patient", "kindly wait", "give me a moment",
            "bear with me", "checking", "under review", "please hold", "one moment"]

CUST_ERRORS = [
    (["abnormal betting", "unusual betting", "irregular betting", "winning amount", "deduct",
      "suspicious bet", "betting pattern"],
     "Withdrawal rejected for abnormal/unusual betting; winnings deducted",
     "Show the customer the evidence and the appeal path; escalate to Risk; apply one consistent deduction SOP."),
    (["deposit", "deposito", "not credited", "not reflect", "hindi pumasok", "hindi naipasok",
      "stuck", "naipit", "pending deposit", "load", "top up", "recharge"],
     "Deposit not credited / stuck (GCash/PSP delay)",
     "Check the PSP/GCash status and the receipt; credit or refund within SLA; give a ticket/ETA instead of 'wait'."),
    (["withdraw", "withdrawal", "payout", "cash out", "cashout", "encash", "widraw"],
     "Withdrawal pending / not received",
     "Verify the account and processing status; give a clear status or ETA."),
    (["otp", "verification code", "sim", "mobile number", "registered number", "unbind",
      "change number", "cannot receive code"],
     "OTP not received / cannot change registered number (e.g., lost SIM)",
     "Apply the lost-SIM / change-number exception flow; don't dead-end the customer."),
    (["bonus", "turnover", "promo", "free spin", "scatter", "rebate", "cashback"],
     "Bonus / turnover / promotion terms confusion",
     "Explain the turnover/promo terms clearly and verify eligibility."),
    (["locked", "frozen", "cannot log", "can't log", "cannot login", "password", "forgot user",
      "account lock"],
     "Account access / locked out",
     "Walk through the verification & unlock steps once, then unlock or escalate; avoid repeating requirements."),
    (["error", "glitch", "maintenance", "lag", "disconnect", "stuck", "server"],
     "Game error / technical glitch",
     "Check the server logs and confirm with the technical team; inform the customer of the outcome."),
]


def detect_error(cust_texts, human_texts, labels):
    c = " ".join(cust_texts).lower()
    for kws, err, sol in CUST_ERRORS:
        if any(k in c for k in kws):
            return err, sol
    if "No reply" in labels:
        return "No human agent ever replied (bot-only)", "Review routing/staffing and call the customer back."
    if "No solution provided" in labels:
        return ("Agent gave only holding replies - no resolution",
                "Escalate to an agent with authority; give an ETA or an actual resolution.")
    if "Cursing/Profanity" in labels:
        return "Customer emotional/abusive (no specific request)", "Assign a senior agent to de-escalate; check account/betting status."
    return "No specific issue detected in the transcript", "Review the transcript and follow up as needed."


def summarize_problem(cust_texts):
    if not cust_texts:
        return ""
    best = max(cust_texts, key=len)
    return best.strip()[:240]

STATUS = {"running": False, "date": None, "group": None, "page": 0, "chats": 0, "error": None, "done": False}
LOCK = threading.Lock()
GROUPS = {}

SECRET = (os.environ.get("SESSION_SECRET") or CFG.get("session_secret") or "change-me").encode()
_AENV_U = os.environ.get("ADMIN_USER")
_AENV_P = os.environ.get("ADMIN_PASSWORD")
if _AENV_U and _AENV_P:
    _salt = os.urandom(16)
    _dk = hashlib.pbkdf2_hmac("sha256", _AENV_P.encode(), _salt, 200000)
    AUTH = {"username": _AENV_U, "salt": _salt.hex(), "hash": _dk.hex(), "iter": 200000}
else:
    AUTH = CFG.get("auth") or {}
COOKIE = "lcmon"
SESSION_HOURS = 12
LOGIN_FAILS = {}
LOGIN_MAX = 8
LOGIN_WINDOW = 900


def check_login(user, pw):
    a = AUTH
    if not a or user != a.get("username"):
        return False
    try:
        dk = hashlib.pbkdf2_hmac("sha256", pw.encode(), bytes.fromhex(a["salt"]), int(a.get("iter", 200000)))
        return hmac.compare_digest(dk.hex(), a["hash"])
    except Exception:
        return False


def make_session(user):
    exp = int(time.time()) + SESSION_HOURS * 3600
    p = base64.urlsafe_b64encode(json.dumps({"u": user, "e": exp}).encode()).decode().rstrip("=")
    sig = hmac.new(SECRET, p.encode(), hashlib.sha256).hexdigest()
    return p + "." + sig


def verify_session(tok):
    try:
        p, sig = tok.split(".", 1)
        if not hmac.compare_digest(hmac.new(SECRET, p.encode(), hashlib.sha256).hexdigest(), sig):
            return None
        data = json.loads(base64.urlsafe_b64decode(p + "=" * (-len(p) % 4)))
        if data["e"] < time.time():
            return None
        return data["u"]
    except Exception:
        return None


LOGIN_HTML = """<!doctype html><html lang="en"><head><meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>LiveChat Bad-Rating Monitor</title>
<link rel="icon" type="image/svg+xml" href="/favicon.svg"/>
<style>
*{box-sizing:border-box}
body{margin:0;height:100vh;display:flex;align-items:center;justify-content:center;font:14px/1.5 system-ui,Segoe UI,Roboto,Arial;color:#e7eefc;
 background:radial-gradient(1200px 600px at 20% -10%,#1b2b52,transparent 60%),linear-gradient(180deg,#0b1426,#0a1120)}
.box{width:340px;padding:26px;border-radius:16px;background:linear-gradient(160deg,rgba(59,130,246,.14),rgba(139,92,246,.06)),#111c33;border:1px solid #2b3f63;box-shadow:0 30px 60px -30px #000}
h1{font-size:17px;margin:0 0 4px}p.sub{color:#8ba0c4;margin:0 0 16px;font-size:12.5px}
label{display:block;color:#8ba0c4;font-size:12px;margin:10px 0 4px}
input{width:100%;padding:10px 12px;border-radius:10px;border:1px solid #22304d;background:#0b1426;color:#e7eefc}
button{width:100%;margin-top:16px;padding:11px;border:0;border-radius:10px;font-weight:700;color:#fff;cursor:pointer;background:linear-gradient(135deg,#3b82f6,#06b6d4)}
.err{color:#ff5d6c;font-size:12.5px;margin-top:10px;min-height:16px}
</style></head><body>
<form class="box" onsubmit="return doLogin(event)">
<h1>LiveChat Bad-Rating Monitor</h1><p class="sub">Sign in to continue</p>
<label>Username</label><input id="u" autocomplete="username" autofocus/>
<label>Password</label><input id="p" type="password" autocomplete="current-password"/>
<button type="submit">Sign in</button><div class="err" id="err">{{MSG}}</div>
</form>
<script>
async function doLogin(e){e.preventDefault();document.getElementById('err').textContent='';
 const r=await fetch('/api/login',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({user:u.value,pass:p.value})});
 if(r.ok){location.href='/';}else{const d=await r.json().catch(()=>({}));document.getElementById('err').textContent=d.error||'Login failed';}
 return false;}
</script></body></html>"""


def post(url, body, tries=6):
    for i in range(tries):
        try:
            r = requests.post(url, headers=HDR, json=body, timeout=(10, 120))
            if r.status_code == 429:
                time.sleep(3 + 2 * i); continue
            if r.status_code >= 500:
                time.sleep(2); continue
            return r
        except Exception:
            time.sleep(2)
    return None


def load_groups():
    global GROUPS
    p = os.path.join(DATA, "groups.json")
    r = post(CFG_API + "list_groups", {"all": True})
    if r is not None and r.status_code == 200:
        GROUPS = {g["id"]: g.get("name") for g in r.json()}
        json.dump(GROUPS, open(p, "w", encoding="utf-8"), ensure_ascii=False)
    elif os.path.exists(p):
        GROUPS = {int(k): v for k, v in json.load(open(p, encoding="utf-8")).items()}
    return GROUPS


def pht(iso):
    try:
        return datetime.datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(PH)
    except Exception:
        return None


def has(text, words):
    t = (text or "").lower()
    return any(w in t for w in words)


def issue_type(concern, labels, cust_texts):
    joined = " ".join(cust_texts).lower()
    c = (concern or "").lower()
    if "otp" in joined or "verification code" in joined or "otp" in c:
        return "OTP"
    m = {"withdrawal": "Withdrawal", "deposit": "Deposit", "promotions": "Promotions",
         "game issue": "Game", "game": "Game", "sign up": "Sign up", "forget password": "Account"}
    for k, v in m.items():
        if k in c:
            return v
    if "Cursing/Profanity" in labels:
        return "Emotional"
    return "Other"


def classify(cust_texts, human_texts, rating):
    labels = []
    joined_c = " ".join(cust_texts).lower()
    if has(joined_c, PROFANITY):
        labels.append("Cursing/Profanity")
    if not human_texts:
        labels.append("No reply")
        if cust_texts:
            labels.append("Abandonment")
    else:
        hold = sum(1 for a in human_texts if any(p in a.lower() for p in HOLD_PAT))
        if hold and hold >= max(1, len(human_texts) * 0.6) and rating == "bad":
            labels.append("No solution provided")
    if has(" ".join(human_texts).lower(), ["cannot help", "not possible", "create a new account",
                                           "cannot unlock", "we cannot", "unable to", "walang magagawa"]) and rating == "bad":
        labels.append("Possible agent mistake (review)")
    if has(joined_c, SIGNAL_WORDS):
        labels.append("Claimed scam/fraud")
    return labels


def suggest(issue, labels, rating):
    s = []
    if "No reply" in labels:
        s.append("No human agent replied - review staffing/routing and follow up with the customer.")
    if "Abandonment" in labels:
        s.append("Customer left unanswered - supervisor to call back.")
    if "No solution provided" in labels:
        s.append("Only holding replies given - escalate to a human with authority / give an ETA or resolution.")
    if "Cursing/Profanity" in labels:
        s.append("Emotional customer - assign a senior agent to de-escalate; check account/betting status with Risk.")
    if "Claimed scam/fraud" in labels:
        s.append("Customer alleges scam - verify the transaction/account and respond with evidence.")
    if "Possible agent mistake (review)" in labels:
        s.append("Possible agent handling error - QA to review this transcript.")
    if issue == "Deposit":
        s.append("Deposit/crediting issue - check PSP/GCash status and credit or refund within SLA.")
    if issue == "Withdrawal":
        s.append("Withdrawal/verification issue - verify and give a clear next step.")
    if issue == "OTP":
        s.append("OTP not received - provide the lost-SIM/change-number exception flow.")
    if not s:
        s.append("Review and follow up as needed.")
    return " ".join(s)


def build_row(c):
    t = c.get("thread", {}) or {}
    users = c.get("users") or []
    cust = next((u for u in users if u.get("type") == "customer"), {}) or {}
    agents = [u for u in users if u.get("type") == "agent" and "@" in (u.get("id") or "")]
    dt = pht(t.get("created_at") or "")
    gids = (t.get("access") or {}).get("group_ids") or []
    gid = gids[-1] if gids else None

    rating_label = None; has_comment = False; comment = ""; concern = ""
    cust_texts = []; human_texts = []; transcript = []; bots = []
    first_c = first_h = None
    for ev in t.get("events", []) or []:
        et = ev.get("type")
        if et == "filled_form":
            for f in ev.get("fields", []) or []:
                lbl = (f.get("label") or ""); ll = lbl.lower()
                ans = f.get("answer")
                val = ans.get("label") if isinstance(ans, dict) else ans
                if "rate your experience" in ll:
                    rating_label = str(val)
                elif "concern" in ll:
                    concern = str(val)
                elif "comment" in ll or "feedback" in ll or "remarks" in ll:
                    if val:
                        has_comment = True; comment = str(val)
            continue
        sm = ev.get("system_message_type")
        if sm == "rating.chat_rated":
            sc = (ev.get("text_vars") or {}).get("score")
            rating_label = rating_label or ("Good" if sc == "good" else "A lot to improve!")
        elif sm == "rating.chat_commented":
            has_comment = True; comment = ev.get("text") or comment
        if et in ("message", "rich_message"):
            a = ev.get("author_id")
            txt = (ev.get("text") or "").replace("\n", " ").strip()
            e2 = pht(ev.get("created_at") or "")
            tstr = e2.strftime("%H:%M:%S") if e2 else ""
            if a == cust.get("id"):
                if txt:
                    cust_texts.append(txt)
                    transcript.append({"role": "customer", "name": cust.get("name") or "Customer", "t": tstr, "text": txt[:500]})
                if first_c is None:
                    first_c = ev.get("created_at")
            else:
                au = next((u for u in users if u.get("id") == a), {})
                nm = au.get("name") or ""
                if "@" in (a or ""):
                    if txt:
                        human_texts.append(txt)
                        transcript.append({"role": "agent", "name": nm or "Agent", "t": tstr, "text": txt[:500]})
                    if first_h is None:
                        first_h = ev.get("created_at")
                elif txt:
                    transcript.append({"role": "bot", "name": nm or "Bot", "t": tstr, "text": txt[:500]})
                    if nm and nm not in bots:
                        bots.append(nm)

    lab = (rating_label or "").lower()
    if lab in ("excellent!", "good", "excellent"):
        rating = "good"
    elif lab in ("could be better", "could be better!", "a lot to improve!", "a lot to improve", "bad"):
        rating = "bad"
    else:
        rating = None

    fr = None
    if first_c and first_h:
        try:
            fr = int((datetime.datetime.fromisoformat(first_h.replace("Z", "+00:00")) -
                      datetime.datetime.fromisoformat(first_c.replace("Z", "+00:00"))).total_seconds())
        except Exception:
            fr = None

    labels = classify(cust_texts, human_texts, rating)
    itype = issue_type(concern, labels, cust_texts)
    err, sol = detect_error(cust_texts, human_texts, labels)
    desc = summarize_problem(cust_texts)
    if comment:
        desc = (desc + "  | comment: " + comment[:160]) if desc else ("comment: " + comment[:160])
    nc = sum(1 for m in transcript if m["role"] == "customer")
    nh = sum(1 for m in transcript if m["role"] == "agent")
    nb = sum(1 for m in transcript if m["role"] == "bot")
    if rating:
        ureason = ""
    elif nc == 0 and nh == 0 and nb == 0:
        ureason = "No messages at all"
    elif nc == 0 and nh == 0 and nb > 0:
        ureason = "Bot only - customer never typed"
    elif nc == 0 and nh > 0:
        ureason = "Agent only - customer never typed"
    elif nh == 0 and nc > 0:
        ureason = "Customer wrote, no human agent replied"
    else:
        ureason = "Agent handled - customer just didn't rate"
    return {
        "chat_id": c.get("id"),
        "time": dt.strftime("%Y-%m-%d %H:%M") if dt else "",
        "group_id": gid, "group_name": GROUPS.get(gid, ""),
        "customer": cust.get("name") or "(unknown)", "customer_id": cust.get("id") or "",
        "rating": rating or "unrated", "rating_label": rating_label or "",
        "has_comment": has_comment, "comment": comment,
        "concern": concern,
        "agents": [{"name": a.get("name"), "id": a.get("id")} for a in agents],
        "n_agents": len(agents),
        "bots": bots,
        "bot_only": len(agents) == 0,
        "issue_type": itype,
        "unrated_reason": ureason,
        "problem": desc,
        "classification": labels,
        "error": err,
        "suggestion": sol,
        "first_response": fr,
        "link": f"https://my.livechatinc.com/archives/{c.get('id')}?query={c.get('id')}",
        "transcript": transcript,
    }


def counts(rows):
    bad = [r for r in rows if r["rating"] == "bad"]
    good = [r for r in rows if r["rating"] == "good"]
    unrated = [r for r in rows if r["rating"] == "unrated"]
    bad_c = [r for r in bad if r["has_comment"]]
    good_c = [r for r in good if r["has_comment"]]
    ur = {}
    for r in unrated:
        k = r.get("unrated_reason") or ""
        if k:
            ur[k] = ur.get(k, 0) + 1
    return {"unrated_reasons": ur, "total": len(rows), "bad": len(bad), "good": len(good), "unrated": len(unrated),
            "rated_commented": len(bad_c) + len(good_c), "bad_commented": len(bad_c),
            "good_commented": len(good_c), "bad_rows": sum(max(1, r["n_agents"]) for r in bad),
            "bad_agent": len([r for r in bad if r["n_agents"] > 0]),
            "bad_bot": len([r for r in bad if r["n_agents"] == 0])}


def fetch_day(date_str, group=None, rated_only=True):
    with LOCK:
        STATUS.update(running=True, date=date_str, group=group, rated=rated_only, page=0, chats=0, error=None, done=False)
    d = datetime.date.fromisoformat(date_str)
    frm = (datetime.datetime(d.year, d.month, d.day, tzinfo=PH).astimezone(datetime.timezone.utc)).strftime("%Y-%m-%dT%H:%M:%S.%f+00:00")
    to = ((datetime.datetime(d.year, d.month, d.day, tzinfo=PH) + datetime.timedelta(days=1)).astimezone(datetime.timezone.utc)).strftime("%Y-%m-%dT%H:%M:%S.%f+00:00")
    filt = {"from": frm, "to": to}
    if group:
        filt["group_ids"] = [int(group)]
    rows = []
    page = None
    while True:
        body = {"page_id": page} if page else {"limit": 100, "sort_order": "asc", "filters": filt}
        r = post(API, body)
        if r is None or r.status_code != 200:
            with LOCK:
                STATUS.update(running=False, error=f"fetch failed at page {STATUS['page']}")
            return
        dd = r.json(); cs = dd.get("chats", [])
        for c in cs:
            rows.append(build_row(c))
        with LOCK:
            STATUS["page"] += 1; STATUS["chats"] += len(cs)
        page = dd.get("next_page_id")
        if not page or not cs:
            break
    if rated_only:
        rows = [r for r in rows if r["rating"] in ("bad", "good")]
    out = {"date": date_str, "group": group, "rated_only": rated_only,
           "generated_at": datetime.datetime.now(PH).strftime("%Y-%m-%d %H:%M:%S"),
           "counts": counts(rows), "rows": rows}
    json.dump(out, open(report_path(date_str, group, rated_only), "w", encoding="utf-8"), ensure_ascii=False)
    with LOCK:
        STATUS.update(running=False, done=True)


def report_path(date_str, group=None, rated_only=True):
    name = f"report_{date_str}"
    if group:
        name += f"_g{group}"
    if rated_only:
        name += "_r"
    return os.path.join(DATA, name + ".json")


_RCACHE = {}


def _read(p):
    if p in _RCACHE:
        return _RCACHE[p]
    if os.path.exists(p):
        r = json.load(open(p, encoding="utf-8"))
        _RCACHE[p] = r
        return r
    return None


def _rated_flag(v):
    return str(v) not in ("0", "false", "False", "")


def load_report(date_str, group=None, rated_only=True):
    r = _read(report_path(date_str, group, rated_only))
    if r is None and group:
        r = _read(report_path(date_str, None, rated_only))
    return r


VIEWS = {
    "all": lambda r: True,
    "bad": lambda r: r["rating"] == "bad",
    "bad_agent": lambda r: r["rating"] == "bad" and r["n_agents"] > 0,
    "bad_bot": lambda r: r["rating"] == "bad" and r["n_agents"] == 0,
    "good": lambda r: r["rating"] == "good",
    "unrated": lambda r: r["rating"] == "unrated",
    "rated_commented": lambda r: r["rating"] in ("bad", "good") and r["has_comment"],
    "bad_commented": lambda r: r["rating"] == "bad" and r["has_comment"],
    "good_commented": lambda r: r["rating"] == "good" and r["has_comment"],
}


def view_rows(rep, view, group=None, itype=None):
    f = VIEWS.get(view, VIEWS["bad"])
    rows = [r for r in rep["rows"] if f(r)]
    if group:
        rows = [r for r in rows if str(r["group_id"]) == str(group)]
    if itype:
        rows = [r for r in rows if r["issue_type"] == itype]
    if view in ("bad", "bad_agent"):
        flat = []
        for r in rows:
            if r["agents"]:
                for a in r["agents"]:
                    flat.append((r, a))
            else:
                flat.append((r, None))
        return flat, True
    return [(r, None) for r in rows], False


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _sec(self, secure=False):
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Security-Policy",
                         "default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; "
                         "script-src 'self' 'unsafe-inline'; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'")

    def _cookie(self):
        for part in (self.headers.get("Cookie") or "").split(";"):
            k, _, v = part.strip().partition("=")
            if k == COOKIE:
                return v
        return None

    def _user(self):
        t = self._cookie()
        return verify_session(t) if t else None

    def _is_https(self):
        return self.headers.get("X-Forwarded-Proto", "").lower() == "https"

    def _json(self, obj, code=200):
        b = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code); self.send_header("Content-Type", "application/json; charset=utf-8")
        self._sec(); self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b)

    def _html(self, html, code=200):
        b = html.encode("utf-8")
        self.send_response(code); self.send_header("Content-Type", "text/html; charset=utf-8")
        self._sec(); self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b)

    def _set_session(self, user):
        tok = make_session(user)
        sc = "; Secure" if self._is_https() else ""
        self.send_header("Set-Cookie", f"{COOKIE}={tok}; HttpOnly; SameSite=Lax; Path=/; Max-Age={SESSION_HOURS*3600}{sc}")

    def _login_page(self, msg=""):
        html = LOGIN_HTML.replace("{{MSG}}", msg)
        b = html.encode("utf-8")
        self.send_response(200); self.send_header("Content-Type", "text/html; charset=utf-8")
        self._sec(); self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b)

    def do_GET(self):
        u = urllib.parse.urlparse(self.path); q = urllib.parse.parse_qs(u.query)
        if u.path in ("/favicon.ico", "/favicon.svg"):
            fp = os.path.join(PUBLIC, "favicon.svg")
            if os.path.exists(fp):
                b = open(fp, "rb").read()
                self.send_response(200); self.send_header("Content-Type", "image/svg+xml")
                self._sec(); self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b); return
            return self.send_error(404)
        if u.path in ("/login", "/login.html"):
            return self._login_page()
        if not self._user():
            if u.path.startswith("/api/"):
                return self._json({"error": "unauthorized"}, 401)
            return self._login_page()
        if u.path in ("/", "/index.html"):
            b = open(os.path.join(PUBLIC, "index.html"), "rb").read()
            self.send_response(200); self.send_header("Content-Type", "text/html; charset=utf-8")
            self._sec(); self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b); return
        if u.path == "/api/groups":
            return self._json(GROUPS)
        if u.path == "/api/status":
            return self._json(STATUS)
        if u.path == "/api/fetch":
            date = q.get("date", [""])[0]; group = q.get("group", [""])[0] or None
            rated = _rated_flag(q.get("rated", ["1"])[0])
            if STATUS["running"]:
                return self._json({"error": "already running", "status": STATUS}, 409)
            threading.Thread(target=fetch_day, args=(date, group, rated), daemon=True).start()
            return self._json({"started": date, "group": group, "rated": rated})
        if u.path == "/api/report":
            date = q.get("date", [""])[0]; rep = load_report(date, q.get("group", [""])[0] or None, _rated_flag(q.get("rated", ["1"])[0]))
            if not rep:
                return self._json({"error": "not found", "date": date}, 404)
            view = q.get("view", ["bad"])[0]
            g = q.get("group", [None])[0]; ty = q.get("type", [None])[0]
            scope = [r for r in rep["rows"] if (not g or str(r["group_id"]) == str(g)) and (not ty or r["issue_type"] == ty)]
            flat, per_agent = view_rows({"rows": scope}, view)
            out = []
            for r, a in flat:
                x = {k: v for k, v in r.items() if k != "transcript"}
                x["agent_name"] = (a["name"] if a else "; ".join(al["name"] or "" for al in r["agents"]))
                x["shared"] = r["n_agents"] > 1
                out.append(x)
            return self._json({"date": date, "view": view, "counts": counts(scope),
                               "generated_at": rep["generated_at"], "row_count": len(out), "rows": out})
        if u.path == "/api/chat":
            date = q.get("date", [""])[0]; cid = q.get("id", [""])[0]
            g = q.get("group", [""])[0] or None
            rated = _rated_flag(q.get("rated", ["1"])[0])
            row = None
            for gg in (g, None):
                for rr in (rated, not rated):
                    rep = load_report(date, gg, rr)
                    if rep:
                        row = next((r for r in rep["rows"] if r["chat_id"] == cid), None)
                        if row:
                            break
                if row:
                    break
            return self._json(row or {"error": "chat not found"}, 200 if row else 404)
        if u.path == "/api/download":
            date = q.get("date", [""])[0]; view = q.get("view", ["bad"])[0]
            rep = load_report(date, q.get("group", [""])[0] or None, _rated_flag(q.get("rated", ["1"])[0]))
            if not rep:
                return self._json({"error": "not found"}, 404)
            return self._download(rep, date, view, q.get("group", [None])[0], q.get("type", [None])[0])
        self.send_error(404)

    def do_POST(self):
        u = urllib.parse.urlparse(self.path)
        n = int(self.headers.get("Content-Length") or 0)
        raw = self.rfile.read(n) if n else b""
        if u.path == "/api/login":
            ip = self.client_address[0]; now = time.time()
            c = LOGIN_FAILS.get(ip)
            if c and c[1] > now and c[0] >= LOGIN_MAX:
                return self._json({"error": "Too many attempts. Try again later."}, 429)
            try:
                data = json.loads(raw or b"{}")
            except Exception:
                data = {}
            if check_login(str(data.get("user", "")), str(data.get("pass", ""))):
                LOGIN_FAILS.pop(ip, None)
                self.send_response(200); self._set_session("Admin")
                self._sec(); self.send_header("Content-Type", "application/json; charset=utf-8")
                b = b'{"ok":true}'; self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b)
                return
            if not c or c[1] < now:
                LOGIN_FAILS[ip] = [1, now + LOGIN_WINDOW]
            else:
                c[0] += 1
            return self._json({"error": "Invalid username or password."}, 401)
        if u.path == "/api/logout":
            self.send_response(200)
            self.send_header("Set-Cookie", f"{COOKIE}=; HttpOnly; SameSite=Lax; Path=/; Max-Age=0")
            self._sec(); self.send_header("Content-Type", "application/json; charset=utf-8")
            b = b'{"ok":true}'; self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b)
            return
        self.send_error(404)

    def _download(self, rep, date, view, group, itype):
        flat, per_agent = view_rows(rep, view, group, itype)
        wb = Workbook(); ws = wb.active; ws.title = view[:30]
        headers = ["Date/Time", "Group Name", "Group ID", "Chat ID", "Open LiveChat", "Customer",
                   "Rating", "Rated Label", "Commented", "Issue Type", "Agents (all)", "Agent (row)", "Bot",
                   "Classification", "Problem Description", "Error / Root cause", "Solution"]
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            ws.cell(1, c).font = Font(bold=True)
        red = PatternFill("solid", fgColor="FFC7CE")
        r = 2
        for row, a in flat:
            vals = [row["time"], row["group_name"], row["group_id"], row["chat_id"], "open", row["customer"],
                    row["rating"], row["rating_label"], "yes" if row["has_comment"] else "", row["issue_type"],
                    "; ".join(al["name"] or "" for al in row["agents"]), (a["name"] if a else ""),
                    "; ".join(row.get("bots") or []),
                    ", ".join(row["classification"]), row["problem"], row.get("error", ""), row["suggestion"]]
            for i, v in enumerate(vals, 1):
                ws.cell(r, i, v)
            url = f"https://my.livechatinc.com/archives/{row['chat_id']}?query={row['chat_id']}"
            cid = ws.cell(r, 4); cid.hyperlink = url; cid.font = Font(color="0563C1", underline="single")
            lc = ws.cell(r, 5); lc.hyperlink = url; lc.font = Font(color="0563C1", underline="single")
            if row["n_agents"] > 1:
                for c in range(1, len(headers) + 1):
                    ws.cell(r, c).fill = red
            r += 1
        for i, w in enumerate([16, 26, 9, 14, 13, 16, 8, 14, 11, 13, 28, 14, 16, 30, 50, 40, 60], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        buf = io.BytesIO(); wb.save(buf); data = buf.getvalue()
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="LiveChat_{view}_{date}.xlsx"')
        self._sec(); self.send_header("Content-Length", str(len(data))); self.end_headers(); self.wfile.write(data)


def scheduler():
    # Auto-fetch yesterday (all groups, rated only) each day so data is ready.
    while True:
        try:
            target = (datetime.datetime.now(PH).date() - datetime.timedelta(days=1)).isoformat()
            if not os.path.exists(report_path(target, None, True)) and not STATUS["running"]:
                print("scheduler: auto-fetching", target)
                fetch_day(target, None, True)
        except Exception as e:
            print("scheduler error:", e)
        time.sleep(900)


if __name__ == "__main__":
    load_groups()
    threading.Thread(target=scheduler, daemon=True).start()
    srv = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"LC Monitor on http://{HOST}:{PORT} groups={len(GROUPS)} auth={'env' if AUTH else 'none'} token={'set' if TOKEN else 'MISSING'}")
    srv.serve_forever()
